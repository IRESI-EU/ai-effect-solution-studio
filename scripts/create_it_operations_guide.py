from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "AI_EFFECT_Solution_Studio_Production_Operations_Guide.docx"

NAVY = "17365D"
BLUE = "2F67D8"
PALE_BLUE = "EEF4FC"
PALE_GRAY = "F6F7F9"
MID_GRAY = "667085"
LIGHT_BORDER = "D9D9D9"
BLACK = "000000"
WHITE = "FFFFFF"
RED = "B42318"
AMBER = "9A6700"
GREEN = "16794B"


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def keep_row_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Aptos", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def shade_run(run, color):
    r_pr = run._element.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    r_pr.append(shd)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def remove_paragraph_borders(paragraph_or_style):
    element = paragraph_or_style._p if hasattr(paragraph_or_style, "_p") else paragraph_or_style.element
    p_pr = element.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def add_table(doc, headers, rows, widths=None, font_size=8.8, alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    repeat_header(header)
    keep_row_together(header)
    for index, text in enumerate(headers):
        cell = header.cells[index]
        set_cell_fill(cell, NAVY)
        set_cell_margins(cell, top=115, bottom=115)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            cell.width = widths[index]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(str(text))
        set_run_font(run, size=8.5, bold=True, color=WHITE)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        keep_row_together(row)
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cell.width = widths[col_index]
            if row_index % 2:
                set_cell_fill(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if alignments and alignments[col_index] == "center":
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(value))
            set_run_font(run, size=font_size, color=BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Cm(0.6 + 0.5 * level)
        p.paragraph_format.first_line_indent = Cm(-0.3)
        if isinstance(item, tuple):
            lead, remainder = item
            r = p.add_run(lead)
            set_run_font(r, size=10, bold=True)
            r = p.add_run(remainder)
            set_run_font(r, size=10)
        else:
            r = p.add_run(item)
            set_run_font(r, size=10)


def add_numbered(doc, items):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.55)
        p.paragraph_format.space_after = Pt(4)
        number = p.add_run(f"{index}.  ")
        set_run_font(number, size=10, bold=True)
        r = p.add_run(item)
        set_run_font(r, size=10)


def add_label_paragraph(doc, label, text, color=BLACK):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(label)
    set_run_font(r, size=10, bold=True, color=color)
    r = p.add_run(text)
    set_run_font(r, size=10)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.12
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_run_font(r, size=10.2, bold=True)
        r = p.add_run(text[len(bold_lead):])
        set_run_font(r, size=10.2)
    else:
        r = p.add_run(text)
        set_run_font(r, size=10.2)
    return p


def page_break(doc):
    doc.add_page_break()


def add_code_block(doc, lines):
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.55)
        p.paragraph_format.right_indent = Cm(0.4)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.keep_together = True
        run = p.add_run(line)
        set_run_font(run, name="Consolas", size=8.5, color=BLACK)
        shade_run(run, PALE_GRAY)


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(7)

    title_style = styles["Title"]
    title_style.font.name = "Aptos Display"
    title_style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title_style.font.size = Pt(27)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor.from_string(BLACK)
    remove_paragraph_borders(title_style)

    for name, size in (("Heading 1", 17), ("Heading 2", 12.5), ("Heading 3", 10.5)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(6)

    footer = section.footer
    p = footer.paragraphs[0]
    p.paragraph_format.tab_stops.add_tab_stop(Cm(8.7), WD_TAB_ALIGNMENT.CENTER)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(17.4), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run("AI EFFECT Solution Studio")
    set_run_font(r, size=7.5, color=MID_GRAY)
    p.add_run("\t")
    r = p.add_run("Internal IT handover")
    set_run_font(r, size=7.5, color=MID_GRAY)
    p.add_run("\t")
    r = p.add_run("Page ")
    set_run_font(r, size=7.5, color=MID_GRAY)
    add_page_number(p)

    # Title page
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run("AI EFFECT")
    set_run_font(r, name="Aptos Display", size=12, bold=True, color=BLUE)
    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_after = Pt(10)
    p.add_run("AI EFFECT Solution Studio Production Operations Guide")
    remove_paragraph_borders(p)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(22)
    r = p.add_run("Infrastructure ownership security cost management and operational procedures")
    set_run_font(r, name="Aptos Display", size=14, color=MID_GRAY)

    add_body(doc, "This guide provides the IT team with the information required to operate, secure, monitor, back up, recover, and budget for the production AI EFFECT Solution Studio and its four country node use cases. It records the verified production state as of 4 October 2026 and identifies the controls that still require implementation.")

    add_table(
        doc,
        ["Document field", "Value"],
        [
            ["Audience", "AI EFFECT IT operations, infrastructure and security teams"],
            ["Environment", "Production"],
            ["Primary service", "https://platform.ai-effect.eu/"],
            ["Secondary hostname", "https://portal.renewenergy.io/"],
            ["Production host", "evpa-uat at 212.47.78.250"],
            ["Hosting provider", "Contabo GmbH"],
            ["Document version", "1.0"],
            ["Status date", "4 October 2026"],
            ["Git branch", "feature/production-ui-polish-v1"],
            ["Recorded commit", "ddcfc3b"],
        ],
        widths=[Cm(5.0), Cm(12.0)],
        font_size=9.2,
    )
    add_label_paragraph(doc, "Classification  ", "Internal operational documentation. Do not attach credentials, passwords, private keys, tokens, or environment-file contents.", color=RED)

    page_break(doc)
    add_heading(doc, "Contents", 1)
    contents = [
        "Executive summary",
        "Production architecture",
        "Infrastructure inventory",
        "Country node services and data",
        "Capacity and storage",
        "Expenses and budget planning",
        "Security controls",
        "Backup and recovery",
        "Monitoring and alerting",
        "Routine operations and change management",
        "Incident response and rollback",
        "Risks and prioritized action plan",
        "Ownership and access management",
        "Appendix A Operational commands",
        "Appendix B Reference sources",
    ]
    add_numbered(doc, contents)
    add_heading(doc, "Document use", 2)
    add_body(doc, "Sections 2 through 6 explain what is running and what it costs. Sections 7 through 11 define the required operating controls. Section 12 should be used as the initial remediation backlog. Appendix A contains read-only commands for routine diagnosis.")

    page_break(doc)
    add_heading(doc, "1 Executive summary", 1)
    add_body(doc, "The AI EFFECT Solution Studio and all four country node use cases currently run as Docker containers on one shared Contabo VPS. The portal does not depend on a developer laptop. Browser traffic reaches Traefik over HTTPS, the Solution Studio communicates with a private API and WP3 orchestrator, and the orchestrator invokes country services over private Docker networks.")
    add_body(doc, "The server has sufficient capacity for the initial controlled release. It provides 8 virtual CPUs, 24 GB RAM and a 200 GB NVMe disk. At the time of measurement, 89 GB was used, 97 GB was available, approximately 4.2 GB of RAM was used across the whole server, and approximately 18 GB of RAM remained available.")
    add_body(doc, "The four use cases passed end-to-end production acceptance. Twelve country-node workflow tasks completed successfully and returned output references. The portal catalogue contains sixteen services, including the twelve accepted German, Portuguese, Danish and Dutch services.")
    add_heading(doc, "Immediate IT priorities", 2)
    add_numbered(doc, [
        "Create and test a nightly backup for the new Solution Studio PostgreSQL database and copy it off the VPS.",
        "Add automatic restart policies and Docker health checks to the German containers and Portuguese ClickHouse.",
        "Configure bounded Docker log rotation so a noisy service cannot fill the system disk.",
        "Repair or stop the unrelated evpa-nginx-1 container, which is continuously restarting.",
        "Implement external availability, disk, backup and container-health alerts before broad public promotion.",
    ])
    add_label_paragraph(doc, "Capacity decision  ", "No immediate VPS upgrade is required. A dedicated AI EFFECT VPS should be considered when the system begins handling sustained public traffic or authoritative user data.", color=GREEN)

    page_break(doc)
    add_heading(doc, "2 Production architecture", 1)
    add_body(doc, "The production design separates the public web entry point from private application and country-node services. Only the reverse proxy should accept public web traffic for AI EFFECT.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(10)
    for line in [
        "User browser over HTTPS",
        "        |",
        "Traefik reverse proxy and TLS termination",
        "        |",
        "Solution Studio web application",
        "        | private Docker network",
        "Solution Studio API and PostgreSQL",
        "        |",
        "WP3 orchestrator Redis and workers",
        "        | private ai-effect-services network",
        "German   Portuguese   Danish   Dutch containers",
    ]:
        r = p.add_run(line + "\n")
        set_run_font(r, name="Consolas", size=9.3, bold=line.startswith(("Traefik", "Solution", "WP3", "German")))
    add_body(doc, "The public portal is available at platform.ai-effect.eu and portal.renewenergy.io. Both hostnames resolve to 212.47.78.250 and use the same production application. Traefik manages TLS certificates and routing.")
    add_heading(doc, "Network trust boundaries", 2)
    add_table(
        doc,
        ["Boundary", "Current design", "IT requirement"],
        [
            ["Internet to Traefik", "HTTPS on ports 80 and 443 with automatic certificate handling", "Monitor certificate renewal and redirect HTTP to HTTPS"],
            ["Traefik to web", "Private Docker network; web also binds to 127.0.0.1:13000", "Do not publish the application port externally"],
            ["Web to API", "Private Docker network", "Keep API database credentials server-side"],
            ["API to WP3", "Orchestrator bound to 127.0.0.1:18000 and private networks", "Do not expose orchestrator or Redis publicly"],
            ["WP3 to country nodes", "External private Docker network named ai-effect-services", "Country services must not publish host ports"],
            ["Administrator access", "SSH key authentication to the renew account", "Use named individual keys and remove departed users promptly"],
        ],
        widths=[Cm(3.5), Cm(6.4), Cm(7.1)],
        font_size=8.4,
    )

    page_break(doc)
    add_heading(doc, "3 Infrastructure inventory", 1)
    add_heading(doc, "Server specification", 2)
    add_table(
        doc,
        ["Item", "Observed value", "Operational note"],
        [
            ["Provider", "Contabo GmbH", "Confirmed from the public IP autonomous system"],
            ["Host name", "evpa-uat", "Shared with PowerPal and legacy EVPA workloads"],
            ["Operating system", "Ubuntu 22.04", "Maintain security updates and supported packages"],
            ["Compute", "8 vCPUs on AMD EPYC", "Shared VPS CPU; country workloads may contend during heavy runs"],
            ["Memory", "23.47 GiB usable", "Approximately 18 GiB available at the status snapshot"],
            ["Disk", "196 GiB formatted ext4", "89 GiB used and 97 GiB free at the status snapshot"],
            ["Docker root", "/var/lib/docker", "Contains images, writable layers, logs and volumes"],
            ["Public IP", "212.47.78.250", "Both production hostnames resolve here"],
        ],
        widths=[Cm(3.3), Cm(5.1), Cm(8.6)],
        font_size=8.8,
    )
    add_heading(doc, "Production paths", 2)
    add_table(
        doc,
        ["Purpose", "Path"],
        [
            ["Solution Studio deployment", "/home/renew/ai-effect-solution-studio-staging"],
            ["German Portuguese and Danish repository", "/home/renew/ai-effect-country-nodes/ai-effect-wp3-main"],
            ["Dutch repository", "/home/renew/ai-effect-country-nodes/ai-effect-wp3-dutch"],
            ["German VILLAS source", "/home/renew/ai-effect-country-nodes/villas-node"],
            ["Production overlays and acceptance tools", "/home/renew/ai-effect-country-nodes/overlays"],
            ["WP3 orchestrator", "/home/renew/ai-effect-wp3"],
            ["Legacy portal retained for rollback", "/home/renew/ai-effect-wp3-portal"],
            ["Docker managed data", "/var/lib/docker"],
        ],
        widths=[Cm(6.2), Cm(10.8)],
        font_size=8.8,
    )
    add_label_paragraph(doc, "Important  ", "Do not edit files directly under /var/lib/docker. Manage images, containers and volumes only through Docker and the approved Compose files.", color=RED)

    page_break(doc)
    add_heading(doc, "4 Country node services and data", 1)
    add_table(
        doc,
        ["Node", "Production containers", "Persistent data", "Acceptance result"],
        [
            ["Germany", "data-provider, villas-chronics, villas-node, output-formatter", "germany-node_shared-data", "PASS: ProvideData, GenerateChronics and FormatOutput"],
            ["Portugal", "tef-data-provision, tef-knowledge-store, pt-synthetic-data, tef-clickhouse", "ClickHouse data and synthetic models", "PASS: ExecuteQuery, ApplyFunction, TrainModel and GenerateData"],
            ["Denmark", "dk-config-assistant, dk-topology-generator, dk-draw-topology, dk-ollama", "Denmark data and Ollama model volumes", "PASS: BuildRunConfig, GenerateTopology and RenderTopologies"],
            ["Netherlands", "nl-grid-synth-service, nl-benchmark-runner", "Dutch synthetic-model volume", "PASS: ConfigureAndSynthesize and RunBenchmark"],
        ],
        widths=[Cm(2.1), Cm(5.6), Cm(4.0), Cm(5.3)],
        font_size=8.1,
    )
    add_heading(doc, "Shared platform containers", 2)
    add_bullets(doc, [
        ("Solution Studio  ", "web, API and PostgreSQL containers."),
        ("WP3 execution  ", "orchestrator API, Redis and three worker containers."),
        ("Ingress  ", "the existing powerpal-traefik container terminates TLS and routes both AI EFFECT hostnames."),
        ("Legacy rollback  ", "ai-effect-nginx is stopped; the legacy backend and database remain temporarily available for controlled rollback."),
    ])
    add_heading(doc, "Data ownership", 2)
    add_body(doc, "Git repositories and pinned image definitions make application code reproducible. PostgreSQL, ClickHouse and any user-generated workflow artifacts are stateful and require backups. Ollama models and most generated test data can be downloaded or rebuilt, but rebuilding them takes time and should not be treated as a substitute for a documented recovery procedure.")

    page_break(doc)
    add_heading(doc, "5 Capacity and storage", 1)
    add_heading(doc, "Current usage", 2)
    add_table(
        doc,
        ["Resource", "Current measurement", "Interpretation"],
        [
            ["Server disk", "89 GiB used of 196 GiB", "48 percent used; 97 GiB available"],
            ["All Docker images", "48.61 GB", "Includes AI EFFECT, PowerPal and older EVPA projects"],
            ["Reclaimable images", "23.63 GB", "Do not remove until rollback requirements are reviewed"],
            ["Docker volumes", "6.20 GB", "Includes databases and other projects"],
            ["Docker build cache", "3.52 GB", "Can be cleaned after release stabilization"],
            ["AI EFFECT image estimate", "Approximately 25 to 26 GB gross", "Large images dominate; shared layers may lower true incremental use"],
            ["AI EFFECT persistent volumes", "Approximately 1.4 GB", "Mostly Ollama and ClickHouse data"],
            ["AI EFFECT operational budget", "40 to 50 GB", "Allows for upgrades, old images, logs and moderate data growth"],
        ],
        widths=[Cm(4.2), Cm(4.5), Cm(8.3)],
        font_size=8.7,
    )
    add_heading(doc, "Largest AI EFFECT components", 2)
    add_table(
        doc,
        ["Component", "Approximate size"],
        [
            ["Dutch benchmark image", "6.97 GB"],
            ["Denmark Ollama image", "5.57 GB"],
            ["Denmark Ollama model volume", "0.94 GB"],
            ["Portuguese synthetic-data image", "4.40 GB"],
            ["German VILLAS image", "3.08 GB"],
            ["Dutch synthesizer image", "0.95 GB"],
            ["Portuguese ClickHouse volume", "0.34 GB"],
        ],
        widths=[Cm(11.5), Cm(5.5)],
        font_size=9.0,
        alignments=["left", "center"],
    )
    add_heading(doc, "Capacity thresholds", 2)
    add_bullets(doc, [
        ("Disk warning  ", "70 percent used. Review images, logs, backups and database growth."),
        ("Disk critical  ", "80 percent used. Stop nonessential builds and add capacity or safely remove confirmed obsolete data."),
        ("Memory warning  ", "80 percent sustained for 15 minutes."),
        ("CPU warning  ", "85 percent sustained for 15 minutes, especially during Dutch benchmarking or Ollama execution."),
        ("Scale decision  ", "Move AI EFFECT to a dedicated VPS before sustained public workloads, sensitive user data, or repeated resource contention."),
    ])

    page_break(doc)
    add_heading(doc, "6 Expenses and budget planning", 1)
    add_body(doc, "All four country nodes run inside one fixed-price VPS. Docker does not charge per container, so the four nodes do not create four separate hosting bills. Cost changes occur only when the VPS plan, storage, backup services, domains, or external services change.")
    add_table(
        doc,
        ["Cost item", "Planning estimate", "Notes"],
        [
            ["Existing Contabo VPS", "Approximately EUR 14 per month", "Matches the older Cloud VPS 30 shape. Confirm the actual invoice, VAT, region and billing term in the Contabo account."],
            ["AI EFFECT incremental compute cost", "EUR 0 while sharing the paid VPS", "AI EFFECT consumes capacity but does not add a per-container fee."],
            ["Optional object storage", "From EUR 2.49 per month for 250 GB", "Recommended for encrypted off-server database backups."],
            ["Comparable dedicated AI EFFECT VPS", "Approximately EUR 14 per month", "Would isolate AI EFFECT from PowerPal and EVPA workloads. Pricing changes over time."],
            ["Storage extension", "Plan-specific", "Not required at present; 97 GiB is available."],
            ["Domains DNS and email", "Account-specific", "Managed and billed separately from Docker containers."],
        ],
        widths=[Cm(4.1), Cm(4.1), Cm(8.8)],
        font_size=8.5,
    )
    add_heading(doc, "Budget recommendation", 2)
    add_body(doc, "For the current controlled release, retain the existing VPS and budget for off-server backup storage. For a production service with sustained external usage, budget for a separate AI EFFECT VPS of similar size plus backup storage. This removes competition with unrelated projects and reduces the security blast radius.")
    add_label_paragraph(doc, "Billing caveat  ", "SSH access cannot reveal the contractual invoice. IT or finance must confirm the active Contabo product, renewal date, VAT treatment, add-ons and domain expenses in the relevant accounts.", color=AMBER)

    page_break(doc)
    add_heading(doc, "7 Security controls", 1)
    add_heading(doc, "Controls already in place", 2)
    add_bullets(doc, [
        "HTTPS and valid TLS certificates are provided by Traefik for both production hostnames.",
        "Country services use private Docker networking and publish no host ports.",
        "The Solution Studio web service binds locally to 127.0.0.1:13000.",
        "The WP3 orchestrator binds locally to 127.0.0.1:18000 and Redis has no public host port.",
        "Production environment files are stored on the server with mode 0600.",
        "SSH access uses individual public keys and the former developer keys were removed.",
        "Management pages and management API routes require HTTP Basic Authentication.",
        "The Portuguese production source excludes development secrets, a private test certificate, cached datasets and cached models.",
    ])
    add_heading(doc, "Required security improvements", 2)
    add_table(
        doc,
        ["Priority", "Action", "Reason"],
        [
            ["Critical", "Back up the new database off-server and test restoration", "The active Solution Studio database is not covered by the current legacy backup job"],
            ["High", "Delete the temporary initial-credentials file after secure transfer and rotate the browser credential", "Reduces credential exposure and prevents indefinite reuse"],
            ["High", "Review host and Contabo firewall rules", "Only ports required by approved projects should be reachable from the internet"],
            ["High", "Add Docker log limits", "Default json-file logging may grow without a bound"],
            ["High", "Move to application user accounts or institutional single sign-on", "Basic Authentication is a temporary shared gate, not an auditable identity system"],
            ["Medium", "Apply Ubuntu and Docker security updates on a scheduled maintenance window", "Reduces exposure to known vulnerabilities"],
            ["Medium", "Add container image vulnerability scanning and dependency review", "Country images include large third-party runtimes"],
            ["Medium", "Separate AI EFFECT from unrelated workloads", "Reduces resource contention and shared-host attack surface"],
        ],
        widths=[Cm(2.0), Cm(7.0), Cm(8.0)],
        font_size=8.3,
        alignments=["center", "left", "left"],
    )
    add_heading(doc, "Secret handling rules", 2)
    add_bullets(doc, [
        "Never place passwords, API keys, private keys or environment files in Git, tickets, email or this document.",
        "Store production credentials in an institutional password manager with named-user access and an access log.",
        "Use separate credentials for production, staging and development.",
        "Rotate credentials after staff departure, suspected disclosure or handover to a new operations owner.",
        "Keep the SSH private key off the server and protect it with device encryption and a passphrase where supported.",
    ])

    page_break(doc)
    add_heading(doc, "8 Backup and recovery", 1)
    add_heading(doc, "Current backup state", 2)
    add_body(doc, "A cron job runs at 02:30 each day and retains fourteen days of compressed PostgreSQL dumps. However, the job executes pg_dump inside ai-effect-postgres, which is the legacy portal database. It does not currently protect ai-effect-solution-studio-postgres-1, the active Solution Studio database.")
    add_table(
        doc,
        ["Data set", "Current state", "Required recovery control"],
        [
            ["Solution Studio PostgreSQL", "Not covered by the observed nightly job", "Nightly pg_dump, 14 to 30 day retention, encrypted off-server copy and quarterly restore test"],
            ["Legacy portal PostgreSQL", "Nightly compressed dump with 14 day retention", "Retain temporarily for rollback, then archive according to an approved retirement date"],
            ["Portuguese ClickHouse", "Approximately 342 MB in a Docker volume", "Back up if data is authoritative; otherwise document the repeatable reconstruction procedure"],
            ["Ollama model volume", "Approximately 941 MB", "Can be downloaded again; record the exact model name and version"],
            ["Country repositories and overlays", "Tracked in Git", "Protect GitHub access and maintain pinned commits and release tags"],
            ["Docker images", "Stored locally and partly reproducible", "Keep immutable image tags or registry copies for each production release"],
            ["TLS and Traefik configuration", "Stored on the shared server", "Include configuration and certificate recovery procedure in infrastructure backup"],
        ],
        widths=[Cm(4.1), Cm(5.2), Cm(7.7)],
        font_size=8.2,
    )
    add_heading(doc, "Minimum recovery objectives", 2)
    add_table(
        doc,
        ["Objective", "Initial target", "Meaning"],
        [
            ["Recovery point objective", "24 hours", "At most one day of database changes may be lost"],
            ["Recovery time objective", "4 hours", "Restore portal service within four hours during supported operations"],
            ["Backup retention", "14 to 30 days", "Keep enough restore points to recover from delayed discovery of corruption"],
            ["Restore validation", "Quarterly and after backup changes", "A backup is not accepted until it has been restored and checked"],
        ],
        widths=[Cm(4.0), Cm(4.0), Cm(9.0)],
        font_size=8.8,
    )
    add_heading(doc, "Restore test record", 2)
    add_body(doc, "For every restore test, record the backup timestamp, operator, target environment, start and finish times, verification result, defects found and corrective actions. Restore tests must use an isolated database or disposable environment and must never overwrite production without an approved incident decision.")

    page_break(doc)
    add_heading(doc, "9 Monitoring and alerting", 1)
    add_table(
        doc,
        ["Monitor", "Check", "Suggested threshold", "Response"],
        [
            ["Public availability", "HTTPS GET to platform.ai-effect.eu", "Two failures within five minutes", "Check Traefik, DNS, web health and certificate status"],
            ["API health", "GET /api/health", "Non-200 response or invalid payload", "Check web proxy, API and PostgreSQL"],
            ["Container health", "Docker health and restart state", "Unhealthy, exited or restart loop", "Review logs and dependency health; restart only after identifying the cause"],
            ["Disk", "Root filesystem used percentage", "Warning 70 percent; critical 80 percent", "Review logs, images, build cache, backups and database growth"],
            ["Memory", "Available memory and swap activity", "More than 80 percent used for 15 minutes", "Identify high-memory workflow; apply limits or scale"],
            ["CPU", "Host and container CPU", "More than 85 percent for 15 minutes", "Check Dutch benchmark, Ollama and concurrent jobs"],
            ["Backups", "Age and exit status of newest backup", "No successful backup within 26 hours", "Treat as a production incident"],
            ["TLS", "Certificate expiry", "Less than 21 days remaining", "Investigate Traefik ACME renewal"],
            ["Workflow reliability", "Failed or stuck workflow count", "Any repeated failure or task beyond agreed runtime", "Inspect orchestrator and country-node logs"],
        ],
        widths=[Cm(3.0), Cm(4.2), Cm(4.6), Cm(5.2)],
        font_size=7.8,
    )
    add_heading(doc, "Logging requirements", 2)
    add_bullets(doc, [
        "Configure per-container log rotation, for example a 10 to 50 MB maximum file size and three retained files.",
        "Do not log credentials, authorization headers, full environment variables or private workflow inputs.",
        "Retain security and administrative events long enough for investigation and governance requirements.",
        "Record deployment version, operator, timestamp and verification outcome for each production change.",
    ])

    page_break(doc)
    add_heading(doc, "10 Routine operations and change management", 1)
    add_heading(doc, "Daily checks", 2)
    add_bullets(doc, [
        "Confirm the portal and API health endpoint respond successfully.",
        "Confirm production containers are running and healthy.",
        "Confirm the newest database backup completed and exists off-server.",
        "Review new critical alerts, restart loops and disk usage changes.",
    ])
    add_heading(doc, "Weekly checks", 2)
    add_bullets(doc, [
        "Review disk, image, volume and log growth.",
        "Review failed and long-running workflows.",
        "Review SSH access and recent production changes.",
        "Confirm DNS, TLS and Basic Authentication behavior on both hostnames.",
    ])
    add_heading(doc, "Monthly checks", 2)
    add_bullets(doc, [
        "Apply approved Ubuntu, Docker and application security updates during a maintenance window.",
        "Review container image vulnerabilities and pinned dependencies.",
        "Review Contabo usage, invoice, domain renewal status and backup storage consumption.",
        "Remove obsolete images only after confirming rollback requirements and successful backups.",
        "Review whether shared-host isolation remains acceptable.",
    ])
    add_heading(doc, "Controlled production update", 2)
    add_numbered(doc, [
        "Record the change request, release commit, affected services, owner, risk and rollback version.",
        "Confirm the Git working tree is clean and the release commit is pushed to the institutional repository.",
        "Create and verify a fresh database backup and, for major updates, a Contabo snapshot.",
        "Validate the Docker Compose configuration and confirm secret files retain mode 0600.",
        "Build or pull immutable images. Avoid relying only on the latest tag for production releases.",
        "Deploy the affected services without deleting persistent volumes.",
        "Verify HTTPS, API health, public catalogue, protected management paths and container health.",
        "Run the four-country acceptance validator when orchestration, service metadata or country images change.",
        "Record the result and retain the previous image set until the agreed stabilization period ends.",
    ])
    add_label_paragraph(doc, "Change rule  ", "Never run docker compose down with volume deletion flags in production unless a separately approved data-destruction procedure identifies the exact volumes and confirms recoverable backups.", color=RED)

    page_break(doc)
    add_heading(doc, "11 Incident response and rollback", 1)
    add_heading(doc, "Initial response", 2)
    add_numbered(doc, [
        "Confirm the incident scope: one browser, one hostname, one container, one country node or the complete server.",
        "Preserve evidence. Record timestamps, recent changes, container status and relevant logs before restarting services.",
        "Protect data. Stop further workflow submissions if they could cause corruption or repeated failed writes.",
        "Restore the smallest affected layer first: container, application release, database or full VPS.",
        "Verify service health and business behavior, not only that a container is running.",
        "Document the cause, recovery steps and preventive action.",
    ])
    add_heading(doc, "Rollback options", 2)
    add_table(
        doc,
        ["Failure", "Preferred rollback", "Caution"],
        [
            ["New web or API image", "Redeploy the previously recorded immutable image tag", "Do not change database state unless migration compatibility is confirmed"],
            ["Country-node image", "Restore the previous pinned image and rerun that country acceptance workflow", "Preserve shared volumes unless corruption is proven"],
            ["Database migration", "Use the documented application rollback or restore a verified pre-change dump", "Database restore can discard newer data and requires incident approval"],
            ["Traefik routing", "Reapply the last known-good labels and verify both hostnames", "Avoid exposing private API, orchestrator or country ports"],
            ["Complete release failure", "Temporarily re-enable the retained legacy frontend only under an approved rollback plan", "The legacy service is not the long-term production target"],
            ["Server loss", "Provision a replacement VPS, restore configuration and databases, then repoint DNS", "Requires off-server backups and documented infrastructure dependencies"],
        ],
        widths=[Cm(3.5), Cm(7.2), Cm(6.3)],
        font_size=8.2,
    )

    page_break(doc)
    add_heading(doc, "12 Risks and prioritized action plan", 1)
    add_table(
        doc,
        ["Priority", "Risk", "Required action", "Target"],
        [
            ["P0", "Active database is not in the observed nightly backup", "Create off-server Solution Studio backups and complete a restore test", "Before accepting important user-created solutions"],
            ["P1", "German services and ClickHouse may not restart after a host reboot", "Apply unless-stopped policies and Docker health checks", "Before next planned reboot"],
            ["P1", "Docker json-file logs have no confirmed bound", "Set daemon or per-service log rotation limits", "Within one week"],
            ["P1", "Old evpa-nginx-1 is in a continuous restart loop", "Identify owner, preserve necessary evidence, then repair or stop it", "Immediate shared-host housekeeping"],
            ["P1", "Basic Authentication uses a shared credential", "Rotate credential and plan named-user authentication or institutional SSO", "Before wider team access"],
            ["P2", "AI EFFECT shares the server with unrelated projects and public ports", "Review firewall and plan dedicated hosting when usage increases", "Before sustained public production"],
            ["P2", "Large images consume disk and slow recovery", "Create CPU-appropriate images where possible and retain immutable release tags", "Next image optimization cycle"],
            ["P2", "No external alerting was verified", "Implement uptime, disk, backup and container-health alerts", "Before public announcement"],
            ["P3", "Legacy portal resources remain on the host", "Set a retirement date after the rollback window and archive required data", "After release stabilization"],
        ],
        widths=[Cm(1.8), Cm(4.4), Cm(6.8), Cm(4.0)],
        font_size=7.8,
        alignments=["center", "left", "left", "left"],
    )
    add_heading(doc, "Release readiness decision", 2)
    add_body(doc, "The platform is technically operational and has passed four-country functional acceptance. IT should treat backup coverage, restart behavior, log rotation and alerting as release-hardening work. The database backup gap is the only item that should prevent storing business-critical user data until corrected and tested.")

    page_break(doc)
    add_heading(doc, "13 Ownership and access management", 1)
    add_table(
        doc,
        ["Responsibility", "Suggested owner", "Record to maintain"],
        [
            ["Contabo account and invoices", "Institutional IT or finance", "Account owner, renewal date, billing contact, plan and add-ons"],
            ["DNS and domains", "Institutional IT", "Registrar, expiry, DNS zone owner and recovery contact"],
            ["SSH administration", "Named IT administrators", "Authorized key owner, issue date, review date and removal date"],
            ["Production deployment", "Application operations owner", "Release commit, images, approver, deployment time and validation result"],
            ["Database backups", "Infrastructure operations", "Backup age, off-server location, retention and restore-test evidence"],
            ["Security response", "Institutional security contact", "Incident severity, evidence, decisions, notifications and remediation"],
            ["Country-node functional ownership", "Relevant node technical leads", "Pinned source version, image version and acceptance result"],
        ],
        widths=[Cm(4.5), Cm(4.5), Cm(8.0)],
        font_size=8.5,
    )
    add_heading(doc, "Access review checklist", 2)
    add_bullets(doc, [
        "Every administrator has an individual SSH key; shared private keys are prohibited.",
        "Departed or transferred staff are removed from authorized_keys and password-manager access.",
        "Production browser credentials are rotated and distributed only through the password manager.",
        "GitHub organization permissions and deployment branch protection are reviewed.",
        "Contabo, DNS and domain accounts use multifactor authentication and institutional recovery contacts.",
        "At least two authorized IT staff can recover the service without relying on one developer laptop.",
    ])

    page_break(doc)
    add_heading(doc, "Appendix A Operational commands", 1)
    add_body(doc, "The following commands are read-only and are suitable for routine diagnosis after connecting to the server as the renew user. They do not display secret values.")
    add_heading(doc, "Portal and API checks", 2)
    add_code_block(doc, [
        "curl -fsS -o /dev/null -w '%{http_code}\\n' https://platform.ai-effect.eu/",
        "curl -fsS https://platform.ai-effect.eu/api/health",
        "curl -fsS -L https://platform.ai-effect.eu/api/services | head",
    ])
    add_heading(doc, "Docker and capacity checks", 2)
    add_code_block(doc, [
        "docker ps --format 'table {{.Names}}\\t{{.Status}}\\t{{.Ports}}'",
        "docker stats --no-stream",
        "docker system df",
        "df -hT /",
        "free -h",
    ])
    add_heading(doc, "Service logs", 2)
    add_code_block(doc, [
        "docker logs --tail 200 ai-effect-solution-studio-web-1",
        "docker logs --tail 200 ai-effect-solution-studio-api-1",
        "docker logs --tail 200 orchestrator-api-1",
        "docker logs --tail 200 <country-container-name>",
    ])
    add_heading(doc, "Four-country acceptance", 2)
    add_code_block(doc, [
        "cd /home/renew/ai-effect-country-nodes/overlays",
        "python3 import_release_candidate.py",
        "python3 validate_release_candidate.py",
    ])
    add_label_paragraph(doc, "Operational caution  ", "Do not run prune, volume deletion, database restore, firewall modification or compose-down commands during diagnosis unless the exact impact and rollback path have been approved.", color=RED)

    page_break(doc)
    add_heading(doc, "Appendix B Reference sources", 1)
    references = [
        ("Production acceptance record", "docs/PRODUCTION_ACCEPTANCE_2026-10-04.md in the Solution Studio repository"),
        ("Previous developer handover", "Handover-to-Adarsh.md supplied to the project"),
        ("Contabo VPS product information", "https://contabo.com/en/vps-server/"),
        ("Contabo Cloud VPS documentation", "https://docs.contabo.com/docs/products/cloud-vps/"),
        ("Contabo Object Storage pricing", "https://contabo.com/en/object-storage/"),
        ("Contabo Auto Backup information", "https://contabo.com/en-us/auto-backup/"),
    ]
    for name, target in references:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(7)
        r = p.add_run(name + "  ")
        set_run_font(r, size=10, bold=True)
        if target.startswith("http"):
            add_hyperlink(p, "View official source", target)
        else:
            r = p.add_run(target)
            set_run_font(r, size=10)
    add_heading(doc, "Measurement notes", 2)
    add_body(doc, "Capacity, container and storage values are point-in-time measurements from 4 October 2026. Docker image size totals may count shared layers differently from physical disk usage. Pricing is a planning estimate based on publicly listed Contabo products and must be reconciled with the actual institutional invoice.")
    add_heading(doc, "Revision history", 2)
    add_table(
        doc,
        ["Version", "Date", "Change"],
        [["1.0", "4 October 2026", "Initial IT operations and production management guide following Solution Studio cutover"]],
        widths=[Cm(2.5), Cm(4.0), Cm(10.5)],
        font_size=9.0,
        alignments=["center", "center", "left"],
    )

    # Improve paragraph widow/orphan control and metadata.
    for paragraph in doc.paragraphs:
        p_pr = paragraph._p.get_or_add_pPr()
        if p_pr.find(qn("w:widowControl")) is None:
            p_pr.append(OxmlElement("w:widowControl"))

    doc.core_properties.title = "AI EFFECT Solution Studio Production Operations Guide"
    doc.core_properties.subject = "Infrastructure security cost management and IT operations"
    doc.core_properties.author = "AI EFFECT Project"
    doc.core_properties.keywords = "AI EFFECT, Solution Studio, production, Docker, IT operations, security"
    doc.core_properties.comments = "Contains no credentials or private key material."

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
